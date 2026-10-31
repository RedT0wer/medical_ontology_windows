"""Главный модуль пайплайна динамического построения онтологии медицинских параметров."""

from typing import Any, Dict, List, Optional
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session
from src.config import logger
from src.db import get_db_session
from src.embeddings import BaseEmbeddingProvider, get_embedding_provider
from src.extractor import EntityExtractor, ExtractedEntity
from src.matcher import MatchDecision, MatchingEngine
from src.models import Parameter, ParameterSynonym, PatientRecord


class QuestionnaireProcessor:
    """Сервис автономной обработки медицинских анкет и динамического расширения онтологии."""

    def __init__(
        self,
        extractor: Optional[EntityExtractor] = None,
        embedding_provider: Optional[BaseEmbeddingProvider] = None,
    ):
        self.extractor = extractor or EntityExtractor()
        self.embedding_provider = embedding_provider or get_embedding_provider()
        logger.info("QuestionnaireProcessor успешно инициализирован.")

    def process_questionnaire(
        self,
        text: str,
        patient_id: str,
        patient_name: Optional[str] = None,
        session: Optional[Session] = None,
    ) -> Dict[str, Any]:
        """Обработка одной медицинской анкеты пациента.
        
        Выполняет:
        1. Извлечение клинических сущностей через LLM (Qwen).
        2. Сопоставление с существующей онтологией (Matching Engine).
        3. Динамическое создание канонических параметров или привязку синонимов.
        4. Сохранение сырых записей пациента в patient_records.
        """
        logger.info(f"--- Начало обработки анкеты пациента {patient_id} ---")

        # 1. Извлечение параметров из текста
        entities: List[ExtractedEntity] = self.extractor.extract_from_text(text, patient_id=patient_id)
        if not entities:
            logger.warning(f"В анкете пациента {patient_id} не обнаружено параметров для сохранения.")
            return {
                "patient_id": patient_id,
                "patient_name": patient_name or "",
                "total_extracted": 0,
                "new_parameters": 0,
                "synonyms_linked": 0,
                "results": [],
            }

        # 2. Выполнение сопоставления и сохранения в рамках сессии
        if session is not None:
            return self._process_entities_in_session(session, entities, patient_id, patient_name)
        else:
            with get_db_session() as sess:
                return self._process_entities_in_session(sess, entities, patient_id, patient_name)

    def _process_entities_in_session(
        self,
        session: Session,
        entities: List[ExtractedEntity],
        patient_id: str,
        patient_name: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Внутренняя обработка списка сущностей внутри активной сессии БД."""
        matcher = MatchingEngine(session=session, embedding_provider=self.embedding_provider)

        new_count = 0
        synonym_count = 0
        processed_details = []

        for entity in entities:
            # Алгоритм сопоставления
            decision: MatchDecision = matcher.match_parameter(
                raw_key=entity.raw_key,
                raw_value=entity.raw_value,
                context=entity.context,
            )

            param_id: Optional[int] = None

            if decision.action == "NEW":
                clean_name = decision.canonical_name.strip()
                # 1. Проверяем среди еще не закоммиченных объектов текущей сессии
                pending_param = next(
                    (
                        obj for obj in session.new 
                        if isinstance(obj, Parameter) and obj.canonical_name.strip().lower() == clean_name.lower()
                    ),
                    None,
                )
                if pending_param:
                    param_id = pending_param.id
                    logger.info(
                        f"Параметр '{clean_name}' уже создан в текущей сессии (ID {param_id}). "
                        f"Привязываем '{entity.raw_key}' как синоним."
                    )
                    synonym_count += 1
                else:
                    # 2. Проверяем в базе данных
                    existing_param = (
                        session.query(Parameter)
                        .filter(Parameter.canonical_name.ilike(clean_name))
                        .first()
                    )

                    if existing_param:
                        param_id = existing_param.id
                        logger.info(
                            f"Параметр '{clean_name}' уже зарегистрирован в базе (ID {param_id}). "
                            f"Привязываем '{entity.raw_key}' как синоним."
                        )
                        synonym_count += 1
                    else:
                        # Создание новой канонической записи
                        param_emb = self.embedding_provider.embed_text(clean_name)
                        meta = {
                            "category": entity.category or "Клинические показатели",
                            "unit": entity.unit or "",
                            "description": decision.reason,
                        }
                        new_param = Parameter(
                            canonical_name=clean_name,
                            description=f"Канонический параметр: {clean_name}",
                            embedding=param_emb,
                            meta_data=meta,
                        )
                        session.add(new_param)
                        session.flush()  # Получаем new_param.id
                        param_id = new_param.id
                        new_count += 1
                        logger.info(f"Создан новый канонический параметр: ID {param_id} ('{new_param.canonical_name}')")

                # Добавляем сырой термин в parameter_synonyms
                self._ensure_synonym(session, param_id, entity.raw_key)

            elif decision.action == "NEW_PARENT_AND_CHILD":
                parent_name = (decision.new_parent_name or "Общий симптом").strip()
                subclass_name = decision.canonical_name.strip()

                # 1. Проверяем или создаем родительский параметр
                parent_param = (
                    session.query(Parameter)
                    .filter(Parameter.canonical_name.ilike(parent_name))
                    .first()
                )
                if not parent_param:
                    for obj in session.new:
                        if isinstance(obj, Parameter) and obj.canonical_name.lower() == parent_name.lower():
                            parent_param = obj
                            break

                if not parent_param:
                    parent_emb = self.embedding_provider.embed_text(parent_name)
                    parent_meta = {
                        "category": entity.category or "Жалобы и симптомы",
                        "unit": "",
                        "description": f"Корневой родительский клинический класс: {parent_name}",
                        "is_parent_class": True,
                    }
                    parent_param = Parameter(
                        canonical_name=parent_name,
                        description=f"Родительский симптом: {parent_name}",
                        embedding=parent_emb,
                        parent_id=None,
                        meta_data=parent_meta,
                    )
                    session.add(parent_param)
                    session.flush()
                    new_count += 1
                    logger.info(f"Создан НОВЫЙ родительский класс: ID {parent_param.id} ('{parent_name}')")

                    # Автоматически связываем с новым родителем другие существующие параметры онтологии,
                    # начинающиеся с этого слова (например, 'Тяжесть в области затылка')
                    related_params = (
                        session.query(Parameter)
                        .filter(
                            Parameter.parent_id.is_(None),
                            Parameter.id != parent_param.id,
                            Parameter.canonical_name.ilike(f"{parent_name}%"),
                        )
                        .all()
                    )
                    for rel in related_params:
                        rel.parent_id = parent_param.id
                        logger.info(
                            f"Существующий параметр ID {rel.id} ('{rel.canonical_name}') "
                            f"автоматически привязан к новому родителю '{parent_name}'"
                        )

                # 2. Создаем или привязываем подкласс
                sub_param = (
                    session.query(Parameter)
                    .filter(Parameter.canonical_name.ilike(subclass_name))
                    .first()
                )
                if not sub_param:
                    for obj in session.new:
                        if isinstance(obj, Parameter) and obj.canonical_name.lower() == subclass_name.lower():
                            sub_param = obj
                            break

                if sub_param:
                    param_id = sub_param.id
                    if not sub_param.parent_id:
                        sub_param.parent_id = parent_param.id
                    synonym_count += 1
                else:
                    sub_emb = self.embedding_provider.embed_text(subclass_name)
                    sub_meta = {
                        "category": entity.category or parent_param.meta_data.get("category", "Клинические показатели"),
                        "unit": entity.unit or "",
                        "description": f"Подкласс параметра '{parent_name}' ({decision.reason})",
                        "is_subclass": True,
                        "parent_id": parent_param.id,
                        "parent_name": parent_name,
                    }
                    sub_param = Parameter(
                        canonical_name=subclass_name,
                        description=f"Подкласс параметра: {parent_name} -> {subclass_name}",
                        embedding=sub_emb,
                        parent_id=parent_param.id,
                        meta_data=sub_meta,
                    )
                    session.add(sub_param)
                    session.flush()
                    param_id = sub_param.id
                    new_count += 1
                    logger.info(
                        f"Создан новый подкласс: ID {param_id} ('{subclass_name}') "
                        f"-> Родитель ID {parent_param.id} ('{parent_name}')"
                    )

                self._ensure_synonym(session, param_id, entity.raw_key)

            elif decision.action == "CHILD_OF":
                clean_name = decision.canonical_name.strip()
                parent_id = decision.parent_id

                # Проверяем в текущей сессии среди еще не сброшенных объектов
                existing_in_session = None
                for obj in session.new:
                    if isinstance(obj, Parameter) and obj.canonical_name.lower() == clean_name.lower():
                        existing_in_session = obj
                        break

                if existing_in_session:
                    param_id = existing_in_session.id
                    if not existing_in_session.parent_id and parent_id:
                        existing_in_session.parent_id = parent_id
                    synonym_count += 1
                else:
                    existing_param = (
                        session.query(Parameter)
                        .filter(Parameter.canonical_name.ilike(clean_name))
                        .first()
                    )
                    if existing_param:
                        param_id = existing_param.id
                        if not existing_param.parent_id and parent_id:
                            existing_param.parent_id = parent_id
                        logger.info(
                            f"Подкласс '{clean_name}' уже зарегистрирован (ID {param_id}). "
                            f"Привязываем '{entity.raw_key}' как синоним."
                        )
                        synonym_count += 1
                    else:
                        param_emb = self.embedding_provider.embed_text(clean_name)
                        parent_param = session.get(Parameter, parent_id) if parent_id else None
                        parent_name = parent_param.canonical_name if parent_param else "общий симптом"
                        meta = {
                            "category": entity.category or (parent_param.meta_data.get("category") if parent_param else "Клинические показатели"),
                            "unit": entity.unit or "",
                            "description": f"Подкласс параметра '{parent_name}' ({decision.reason})",
                            "is_subclass": True,
                            "parent_id": parent_id,
                            "parent_name": parent_name,
                        }
                        new_param = Parameter(
                            canonical_name=clean_name,
                            description=f"Подкласс параметра: {parent_name} -> {clean_name}",
                            embedding=param_emb,
                            parent_id=parent_id,
                            meta_data=meta,
                        )
                        session.add(new_param)
                        session.flush()
                        param_id = new_param.id
                        new_count += 1
                        logger.info(
                            f"Создан новый подкласс онтологии: ID {param_id} ('{clean_name}') "
                            f"-> Родитель ID {parent_id} ('{parent_name}'). Причина: {decision.reason}"
                        )

                self._ensure_synonym(session, param_id, entity.raw_key)

            elif decision.action == "SYNONYM":
                param_id = decision.parameter_id
                synonym_count += 1
                logger.info(
                    f"Параметр '{entity.raw_key}' привязан к существующему ID {param_id} "
                    f"('{decision.canonical_name}'). Причина: {decision.reason}"
                )
                # Добавляем синоним, если такого текста еще нет
                self._ensure_synonym(session, param_id, entity.raw_key)

            # Сохранение записи пациента в patient_records
            record_meta = {
                "patient_name": patient_name or "",
                "context": entity.context,
                "confidence": entity.confidence,
                "category": entity.category,
                "unit": entity.unit,
                "arbitration_reason": decision.reason,
                "similarity": decision.similarity,
            }
            patient_record = PatientRecord(
                patient_id=str(patient_id),
                patient_name=patient_name or None,
                parameter_id=param_id,
                raw_key=entity.raw_key,
                raw_value=entity.raw_value,
                meta_data=record_meta,
            )
            session.add(patient_record)

            processed_details.append(
                {
                    "raw_key": entity.raw_key,
                    "raw_value": entity.raw_value,
                    "action": decision.action,
                    "parameter_id": param_id,
                    "canonical_name": decision.canonical_name,
                    "similarity": round(decision.similarity, 4),
                    "reason": decision.reason,
                }
            )

        logger.info(
            f"[ИНФО] Анкета {patient_id} обработана: создано {new_count} канонических параметров, "
            f"привязано {synonym_count} синонимов (всего параметров: {len(entities)})."
        )

        return {
            "patient_id": patient_id,
            "patient_name": patient_name or "",
            "total_extracted": len(entities),
            "new_parameters": new_count,
            "synonyms_linked": synonym_count,
            "results": processed_details,
        }

    def _ensure_synonym(self, session: Session, parameter_id: int, synonym_text: str) -> None:
        """Добавление синонима в таблицу parameter_synonyms при его отсутствии (безопасно к дубликатам)."""
        clean_text = synonym_text.strip()
        if not clean_text or not parameter_id:
            return

        # 1. Проверяем среди еще не сброшенных объектов текущей сессии
        for obj in session.new:
            if isinstance(obj, ParameterSynonym):
                if obj.parameter_id == parameter_id and obj.synonym_text.strip().lower() == clean_text.lower():
                    return

        # 2. Проверяем в базе данных
        exists = (
            session.query(ParameterSynonym)
            .filter(
                ParameterSynonym.parameter_id == parameter_id,
                ParameterSynonym.synonym_text.ilike(clean_text),
            )
            .first()
        )
        if exists:
            return

        # 3. Атомарная вставка с ON CONFLICT DO NOTHING (защита на уровне СУБД)
        syn_emb = self.embedding_provider.embed_text(clean_text)
        stmt = (
            pg_insert(ParameterSynonym)
            .values(
                parameter_id=parameter_id,
                synonym_text=clean_text,
                embedding=syn_emb,
            )
            .on_conflict_do_nothing(
                constraint="uq_param_synonym"
            )
        )
        session.execute(stmt)
        logger.info(f"Добавлен новый синоним '{clean_text}' для параметра ID {parameter_id}.")

    def process_excel_file(
        self,
        file_path: str = "input.xlsm",
        sheet_name: str = "Лист1",
        limit: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """Чтение и обработка медицинских анкет пациентов напрямую из файла Excel (.xlsm / .xlsx)."""
        import openpyxl

        limit_info = f"{limit} пациентов" if limit and limit > 0 else "все анкеты файла"
        logger.info(f"Загрузка клинических данных из файла {file_path} (лист: '{sheet_name}', объем: {limit_info})...")
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Лист '{sheet_name}' не найден в файле {file_path}. Доступные листы: {wb.sheetnames}")

        sheet = wb[sheet_name]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [cell for cell in next(rows_iter)]

        results = []
        patient_count = 0

        for row in rows_iter:
            if not any(row):
                continue

            row_dict = dict(zip(header, row))
            nom_ib = row_dict.get("nom_ib") or row_dict.get("nom_pp") or str(patient_count + 1)
            patient_id = f"ИБ-{nom_ib}"

            # Формирование связного клинического текста анкеты
            name = str(row_dict.get("name") or "").strip()
            age = row_dict.get("bo3pact")
            gender = row_dict.get("sex")

            text_parts = []
            if name:
                desc = f"Пациент: {name}"
                if age:
                    desc += f", возраст: {age} лет"
                if gender:
                    desc += f", пол: {gender}"
                text_parts.append(desc + ".")

            # Клинические показатели при поступлении
            vitals = []
            if row_dict.get("ad"):
                vitals.append(f"АД при поступлении {str(row_dict['ad']).strip()} мм рт. ст.")
            if row_dict.get("ps"):
                vitals.append(f"пульс {str(row_dict['ps']).strip()} уд/мин")
            if row_dict.get("chdd"):
                vitals.append(f"ЧДД {str(row_dict['chdd']).strip()} в мин")
            if row_dict.get("hb_post"):
                vitals.append(f"гемоглобин {str(row_dict['hb_post']).strip()} г/л")
            if row_dict.get("er_post"):
                vitals.append(f"эритроциты {str(row_dict['er_post']).strip()}")

            if vitals:
                text_parts.append("Витальные показатели при поступлении: " + ", ".join(vitals) + ".")

            jaloba = row_dict.get("jaloba")
            if jaloba and str(jaloba).strip():
                text_parts.append(f"Жалобы и анамнез: {str(jaloba).strip()}")

            full_questionnaire_text = " ".join(text_parts)

            logger.info(f"▶ Обработка пациента {patient_id} ({name or 'Не указан'}) из Excel...")
            res = self.process_questionnaire(
                text=full_questionnaire_text,
                patient_id=patient_id,
                patient_name=name,
            )
            results.append(
                {
                    "patient_id": patient_id,
                    "name": name,
                    "text": full_questionnaire_text,
                    "summary": res,
                }
            )

            patient_count += 1
            if limit and limit > 0 and patient_count >= limit:
                break

        wb.close()
        logger.info(f"Обработка выборки из Excel завершена: обработано пациентов: {len(results)}.")
        return results

    def get_ontology_summary(self) -> Dict[str, Any]:
        """Получение текущей сводки состояния базы знаний (онтологии) с учетом иерархии классов."""
        with get_db_session() as session:
            param_count = session.query(Parameter).count()
            subclass_count = session.query(Parameter).filter(Parameter.parent_id.isnot(None)).count()
            syn_count = session.query(ParameterSynonym).count()
            rec_count = session.query(PatientRecord).count()

            params = session.query(Parameter).all()
            ontology_tree = []
            for p in params:
                synonyms = [s.synonym_text for s in p.synonyms]
                records_cnt = len(p.records)
                children_list = [
                    {"id": c.id, "canonical_name": c.canonical_name, "records_count": len(c.records)}
                    for c in p.children
                ]
                ontology_tree.append(
                    {
                        "id": p.id,
                        "canonical_name": p.canonical_name,
                        "category": p.meta_data.get("category", "Не указана"),
                        "unit": p.meta_data.get("unit", ""),
                        "parent_id": p.parent_id,
                        "parent_name": p.parent.canonical_name if p.parent else None,
                        "children": children_list,
                        "children_count": len(children_list),
                        "synonyms_count": len(synonyms),
                        "synonyms": synonyms,
                        "patient_records_count": records_cnt,
                    }
                )

            return {
                "total_parameters": param_count,
                "total_subclasses": subclass_count,
                "total_synonyms": syn_count,
                "total_patient_records": rec_count,
                "parameters": ontology_tree,
            }

